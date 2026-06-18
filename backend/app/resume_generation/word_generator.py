from io import BytesIO
from typing import Any

from fastapi import HTTPException, status


def _as_list(value: Any) -> list:
    return value if isinstance(value, list) else []


def _format_item(item: Any) -> str:
    if isinstance(item, dict):
        return " | ".join(f"{key.replace('_', ' ').title()}: {value}" for key, value in item.items() if value)
    return str(item)


def generate_resume_docx(content: dict[str, Any]) -> bytes:
    try:
        from docx import Document
    except ImportError as exc:
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail="python-docx is not installed. Run pip install -r requirements.txt.",
        ) from exc

    document = Document()
    document.add_heading(content.get("full_name") or "Resume", 0)
    contact = " | ".join(filter(None, [content.get("email"), content.get("phone")]))
    if contact:
        document.add_paragraph(contact)

    def add_section(title: str) -> None:
        document.add_heading(title, level=1)

    add_section("Professional Summary")
    document.add_paragraph(content.get("professional_summary") or "Not provided")

    add_section("Skills")
    document.add_paragraph(", ".join(str(item) for item in _as_list(content.get("skills"))) or "Not provided")

    add_section("Experience")
    for item in _as_list(content.get("experience")):
        if isinstance(item, dict):
            document.add_paragraph(
                f"{item.get('job_title', '')} - {item.get('company_name', '')} ({item.get('duration', '')})",
                style="List Bullet",
            )
            for point in _as_list(item.get("responsibilities")):
                document.add_paragraph(_format_item(point), style="List Bullet 2")
        else:
            document.add_paragraph(_format_item(item), style="List Bullet")

    for key, title in [("education", "Education"), ("projects", "Projects"), ("certifications", "Certifications")]:
        add_section(title)
        items = _as_list(content.get(key))
        if not items:
            document.add_paragraph("Not provided")
        for item in items:
            document.add_paragraph(_format_item(item), style="List Bullet")

    buffer = BytesIO()
    try:
        document.save(buffer)
    except Exception as exc:
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail=f"Word generation failed: {exc}",
        ) from exc
    buffer.seek(0)
    return buffer.read()
